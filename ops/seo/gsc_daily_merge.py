# -*- coding: utf-8 -*-
"""Daily merge: write GSC per-URL verdicts into the Coverage workbook.

Rules
  * rows that have a real crawl date keep the date in column G; their verdict or
    bucket evidence goes to column H;
  * rows without a crawl date get 网址已收录到 Google / 暂时未被收录 in G
    (+ evidence in H);
  * still-unknown rows stay blank in G and are noted in H;
  * the 汇总 sheet's live counters are refreshed in place (never duplicated).
Safe to run repeatedly.
"""
import datetime
import json
import os
import shutil
import time
import openpyxl
from openpyxl.styles import Font, Alignment

# ── 统一配置 / 干跑层（模式见同目录 preflight.py）──────────────────────
from preflight import preflight, dry_run_enabled, write_path, stage_note, require_env, optional_env



preflight()   # 缺必需变量即报出并退出（不回落默认值）
D = optional_env('DSH_DATA_DIR', require_env('DSH_WORKSPACE', '.wx-bridge 目录（含 plans/）'))
XLSX = require_env('DSH_COVERAGE_XLSX', 'GSC 覆盖率工作簿（写入目标，.xlsx）')
BAK = require_env('DSH_COVERAGE_BAK', '覆盖率工作簿的原始备份（读取来源，.xlsx）')
RESULT = os.path.join(D, "_gsc_results.jsonl")
BUCKETS = os.path.join(D, "_gsc_buckets.json")

INDEXED_TXT = "网址已收录到 Google"
NOT_INDEXED_TXT = "暂时未被收录"

checked = {}
if os.path.exists(RESULT):
    for line in open(RESULT, encoding="utf-8-sig"):
        line = line.strip()
        if not line:
            continue
        try:
            r = json.loads(line)
        except Exception:
            continue
        if r.get("status") in ("indexed", "not_indexed"):
            checked[r["url"].rstrip("/")] = r

bucket_of = {}
if os.path.exists(BUCKETS):
    for reason, v in json.load(open(BUCKETS, encoding="utf-8")).items():
        for u in v["urls"]:
            bucket_of.setdefault(u.rstrip("/"), []).append(f"{reason}（{v['count']}）")

# crawl dates of the 1000 sample, from the untouched original export
dates = {}
if os.path.exists(BAK):
    src = openpyxl.load_workbook(BAK, data_only=True)["表格"]
    for r in range(2, src.max_row + 1):
        u = src.cell(r, 1).value
        if u:
            dates[u.strip().rstrip("/")] = src.cell(r, 2).value

wb = openpyxl.load_workbook(XLSX)
ws = wb["3278全量清单"]
stat = {"dated": 0, "indexed": 0, "not_indexed": 0, "pending": 0}
for r in range(2, ws.max_row + 1):
    url = (ws.cell(r, 2).value or "").rstrip("/")
    if not url:
        continue
    g = ws.cell(r, 7).value
    rec = checked.get(url)
    bucket = bucket_of.get(url)
    if url in dates:                                  # crawled row: date stays in G
        ws.cell(r, 7).value = dates[url]
        ws.cell(r, 7).number_format = "yyyy-mm-dd"
        note = ""
        if bucket:
            note = f"{NOT_INDEXED_TXT}｜GSC 未编入索引原因：" + "；".join(bucket)
        if rec:
            verdict = INDEXED_TXT if rec["status"] == "indexed" else NOT_INDEXED_TXT
            note = f"{verdict}｜{rec.get('state', '')}（{rec['checkedAt'][:16].replace('T', ' ')} 逐条查证）"
        if note:
            ws.cell(r, 8).value = note
        stat["dated"] += 1
        continue
    if isinstance(g, datetime.datetime):              # stray date without source entry
        stat["dated"] += 1
        continue
    if rec:
        ws.cell(r, 7).value = INDEXED_TXT if rec["status"] == "indexed" else NOT_INDEXED_TXT
        ws.cell(r, 7).number_format = "General"
        ws.cell(r, 8).value = f"{rec.get('state', '')}（{rec['checkedAt'][:16].replace('T', ' ')} 逐条查证）"
        stat["indexed" if rec["status"] == "indexed" else "not_indexed"] += 1
        continue
    if bucket:
        ws.cell(r, 7).value = NOT_INDEXED_TXT
        ws.cell(r, 7).number_format = "General"
        ws.cell(r, 8).value = "GSC 未编入索引原因：" + "；".join(bucket)
        stat["not_indexed"] += 1
        continue
    ws.cell(r, 7).value = None
    ws.cell(r, 7).number_format = "General"
    ws.cell(r, 8).value = ("待检查（未出现在 GSC 12 个未编入索引原因中，推定已收录；"
                           "受每日配额限制，按 07:25 定时任务逐日核验）")
    stat["pending"] += 1

# --- refresh live counters in 汇总 ------------------------------------------
ws2 = wb["汇总"]
marker = None
for r in range(1, ws2.max_row + 1):
    v = ws2.cell(r, 1).value
    if isinstance(v, str) and v.startswith("GSC 网址收录检查"):
        marker = r
        break
if marker is None:
    marker = ws2.max_row + 2
    ws2.cell(marker, 1, "GSC 网址收录检查").font = Font(bold=True)
else:
    ws2.cell(marker, 1).value = "GSC 网址收录检查"
    ws2.cell(marker, 1).font = Font(bold=True)
rows = [
    ("已确认收录到 Google", stat["indexed"], "本次逐条网检查证"),
    ("确认未收录", stat["not_indexed"], "逐条查证或命中 GSC“未编入索引”原因清单"),
    ("待逐条核验", stat["pending"], "每日 07:25 自动执行；GSC 每日配额耗尽后次日继续"),
    ("导出清单中已有抓取日期（视为已收录）", stat["dated"], "GSC 已编入索引导出样本"),
]
for i, (a, b, c) in enumerate(rows, start=marker + 1):
    ws2.cell(i, 1, a)
    ws2.cell(i, 2, b)
    ws2.cell(i, 3, c)

# --- save with retry (the file may be open in Excel) ------------------------
for attempt in range(1, 6):
    try:
        wb.save(write_path(XLSX))
        print(f"merge ok: indexed={stat['indexed']} not_indexed={stat['not_indexed']} "
              f"pending={stat['pending']} dated={stat['dated']} (attempt {attempt})")
        break
    except PermissionError:
        print(f"merge: file locked (attempt {attempt}), retrying in 60s")
        time.sleep(60)
else:
    alt = XLSX.replace(".xlsx", f".locked-{datetime.datetime.now():%Y%m%d-%H%M}.xlsx")
    wb.save(write_path(alt))
    print("merge: original locked after 5 tries, wrote", alt)
